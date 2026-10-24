import io
import re
import docx
import requests
from fastapi import FastAPI, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
import wikipedia
from duckduckgo_search import DDGS
import arxiv
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from deep_translator import GoogleTranslator
from pythainlp.tokenize import word_tokenize
app = FastAPI()
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])
translator_to_en = GoogleTranslator(source='auto', target='en')
translator_to_th = GoogleTranslator(source='auto', target='th')
def safe_translate(text, target='th'):
    if not text or len(text.strip()) < 5: return text
    try:
        clean_text = text[:4000]
        return translator_to_en.translate(clean_text) if target == 'en' else translator_to_th.translate(clean_text)
    except: return text
def shorten_text(text, max_len=70):
    if len(text) > max_len:
        return text[:max_len].strip() + "..."
    return text
def add_chapter_two_style(doc, topic, collected_data):
    style = doc.styles['Normal']
    style.font.name = 'TH Sarabun New'
    style.font.size = Pt(16)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    section = doc.sections[0]
    section.top_margin = section.left_margin = Cm(3.81)
    section.bottom_margin = section.right_margin = Cm(2.54)
    h1_1 = doc.add_paragraph()
    h1_1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1_1 = h1_1.add_run("บทที่ 2")
    run1_1.font.size = Pt(18)
    run1_1.font.bold = True
    h1_2 = doc.add_paragraph()
    h1_2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1_2 = h1_2.add_run("แนวคิด ทฤษฎี เอกสาร และงานที่เกี่ยวข้อง")
    run1_2.font.size = Pt(18)
    run1_2.font.bold = True
    doc.add_paragraph() 
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.first_line_indent = Cm(1.52) 
    p_intro.alignment = 9 
    p_intro.add_run(f"ในการศึกษาเรื่อง {topic} ผู้จัดทำได้รวบรวมแนวคิด ทฤษฎีและหลักการ ต่าง ๆ จากเอกสาร และงานที่เกี่ยวข้องดังต่อไปนี้")
    for index, data in enumerate(collected_data, start=1):
        p_idx = doc.add_paragraph()
        p_idx.alignment = WD_ALIGN_PARAGRAPH.LEFT 
        p_idx.paragraph_format.left_indent = Cm(1.52) 
        run_idx = p_idx.add_run(f"2.{index} {data['title']}")
        run_idx.font.bold = True
    doc.add_paragraph() 
    image_counter = 1
    for index, data in enumerate(collected_data, start=1):
        h2 = doc.add_paragraph()
        h2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run2 = h2.add_run(f"2.{index} {data['title']}")
        run2.font.bold = True
        for img_url in data.get('image_urls', []):
            try:
                headers = {
                    'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/14.0 Safari/605.1.15'
                }
                response = requests.get(img_url, headers=headers, timeout=5)
                if response.status_code == 200:
                    image_stream = io.BytesIO(response.content)
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run_img = p_img.add_run()
                    run_img.add_picture(image_stream, width=Cm(10)) 
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    clean_title = data['title'].replace("ข้อมูลสารานุกรมเรื่อง ", "").replace("งานวิจัยเรื่อง ", "")
                    run_cap = p_cap.add_run(f"รูปที่ 2.{image_counter} ภาพประกอบ{clean_title}")
                    run_cap.font.bold = True
                    image_counter += 1
                    doc.add_paragraph() 
            except Exception as e:
                pass 
        blocks = re.split(r'\n\s*\n', data['content'])
        for block in blocks:
            clean_text = block.replace('\n', ' ').strip()
            clean_text = re.sub(r'\s+', ' ', clean_text) 
            if len(clean_text) < 5: continue 
            tokens = word_tokenize(clean_text, engine="newmm")
            spaced_text = '\u200b'.join(tokens)
            p_content = doc.add_paragraph()
            p_content.paragraph_format.first_line_indent = Cm(1.52) 
            p_content.alignment = 9 
            p_content.add_run(spaced_text)
        doc.add_paragraph() 
@app.post("/scout-to-word")
async def scout_to_word(topic: str = Form(...), page_count: int = Form(...)):
    eng_topic = safe_translate(topic, target='en')
    collected_data = []
    image_pool = []
    try:
        with DDGS() as ddgs:
            search_terms = [eng_topic, f"{eng_topic} diagram", f"{eng_topic} technology"]
            for q in search_terms:
                img_results = list(ddgs.images(q, max_results=10))
                for img in img_results:
                    url = img.get('image')
                    if url and url not in image_pool:
                        image_pool.append(url)
    except: pass
    for lang, q in [('en', eng_topic), ('th', topic)]:
        try:
            wikipedia.set_lang(lang)
            p = wikipedia.page(q)
            for img in p.images:
                img_lower = img.lower()
                if img_lower.endswith(('.png', '.jpg', '.jpeg')) and not any(x in img_lower for x in ['icon', 'logo', 'symbol', 'svg']):
                    if img not in image_pool:
                        image_pool.append(img)
        except: pass
    img_idx = 0
    def get_images(count=2):
        nonlocal img_idx
        selected = []
        if not image_pool: return selected 
        attempts = 0
        while len(selected) < count and attempts < len(image_pool):
            img = image_pool[img_idx % len(image_pool)]
            if img not in selected: 
                selected.append(img)
            img_idx += 1
            attempts += 1
        return selected
    for lang, q in [('th', topic), ('en', eng_topic)]:
        try:
            title = f"ข้อมูลสารานุกรมเรื่อง {topic} ({lang.upper()})"
            if len(title) <= 70:
                wikipedia.set_lang(lang)
                p = wikipedia.page(q)
                text = p.content 
                if lang == 'en': text = safe_translate(text[:4000], target='th')               
                collected_data.append({
                    "title": title, 
                    "content": text,
                    "image_urls": get_images(2) 
                })
        except: pass
    try:
        with DDGS() as ddgs:
            search_queries = [
                topic,
                f"ทฤษฎี และหลักการ {topic}",
                f"ประเภท และการใช้งาน {topic}"
            ]
            for q in search_queries:
                results = list(ddgs.text(q, max_results=8)) 
                added_count = 0
                for r in results:
                    if added_count >= 3: 
                        break
                    title = r['title']
                    if len(title) <= 70 and not any(d['title'] == title for d in collected_data):
                        collected_data.append({
                            "title": title, 
                            "content": r['body'],
                            "image_urls": get_images(2) 
                        })
                        added_count += 1
    except: pass
    try:
        client = arxiv.Client()
        search = arxiv.Search(query=eng_topic, max_results=10) 
        added_count = 0
        for r in client.results(search):
            if added_count >= 5: 
                break
            th_title = safe_translate(r.title)
            title = f"งานวิจัยเรื่อง {th_title}"
            if len(title) <= 70:
                th_summary = safe_translate(r.summary)
                collected_data.append({
                    "title": title, 
                    "content": f"บทคัดย่อ: {th_summary}",
                    "image_urls": get_images(2) 
                })
                added_count += 1
    except: pass
    if not collected_data:
        collected_data.append({
            "title": f"ข้อมูลพื้นฐานของ {topic}", 
            "content": f"อยู่ระหว่างการรวบรวมข้อมูล กรุณาลองใช้คีย์เวิร์ดอื่น",
            "image_urls": []
        })
    doc = docx.Document()
    add_chapter_two_style(doc, topic, collected_data)
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return StreamingResponse(
        file_stream, 
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=Chapter_2_Document.docx"}
    )
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
app.mount("/", StaticFiles(directory="dist", html=True), name="static")